#!/usr/bin/env python3
"""
check_coherence.py

Cross-checks a folder of drawings and documents for one product:

  Drawing PDFs (read with the shared sheetreader - text layer or OCR, no
  model needed):
    * the same drawing number on several sheets with a different rev / date / title
    * BOM lines that reference a drawing: is it in the folder, does its title
      match the BOM description, does its NEXT ASSY. point back?
    * NEXT ASSY. links in both directions (a parent BOM must list the child;
      "E-03 A&B" covers BOM refs "E-03A" and "E-03B")
    * drawings printed on the same template form with different general
      tolerances
  Dimension checklists (.docx / .doc, matched to their drawing by file name):
    * material and weight vs the drawing's title block
    * every tolerance row vs the drawing (dimcheck in verify mode)
  Other Word / Excel files (optional, uses Gemma through Ollama):
    * the same part number with different names / revisions / dates

Usage:
    python check_coherence.py FOLDER [--out report.md] [--model gemma4:e4b] [--no-model]
"""

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "dimcheck"))
from sheetreader.normalize import alnum, ids_from_filename, ocr_key, rev_key  # noqa: E402

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


OLLAMA_URL = "http://localhost:11434/api/generate"
DEFAULT_MODEL = "gemma4:e4b"

# ---------------------------------------------------------------------------
# Prompts
# ---------------------------------------------------------------------------

GENERIC_EXTRACTION_PROMPT = """You are extracting structured data from a document excerpt (Word or Excel content).

Find every instance of:
- dates (any format)
- part numbers (alphanumeric codes identifying a part; may include dashes, dots, slashes, spaces)
- part names / descriptions
- revision codes (e.g. "Rev A", "Rev. 3", "-R2", "Revision 1.0")

Return ONLY a JSON array (no other text, no markdown fences). Each element:
{{
  "date": "<date as found, or null>",
  "part_number": "<part number as found, or null>",
  "part_name": "<part name as found, or null>",
  "revision": "<revision as found, or null>",
  "context": "<short snippet of surrounding text, max 15 words>"
}}

Only include an object if at least one field is non-null. Do not invent data.

TEXT EXCERPT:
---
{chunk}
---

JSON array:"""

FUZZY_MATCH_PROMPT = """You are comparing two engineering records to decide if they refer to the
SAME real-world part, despite possibly different wording/formatting.

Record A: {a}
Record B: {b}

Answer with ONLY one word: SAME, DIFFERENT, or UNSURE."""


# ---------------------------------------------------------------------------
# Ollama helpers
# ---------------------------------------------------------------------------

def call_ollama(prompt, model, timeout=180):
    resp = requests.post(
        OLLAMA_URL,
        json={"model": model, "prompt": prompt, "stream": False, "options": {"temperature": 0}},
        timeout=timeout,
    )
    resp.raise_for_status()
    return resp.json().get("response", "")


def extract_json_blob(text, want_array=True):
    text = text.strip()
    text = re.sub(r"^```(json)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()
    pattern = r"\[.*\]" if want_array else r"\{.*\}"
    match = re.search(pattern, text, re.DOTALL)
    if not match:
        return [] if want_array else {}
    try:
        return json.loads(match.group(0))
    except json.JSONDecodeError:
        return [] if want_array else {}


def chunk_text(text, max_chars=3000):
    paras = text.split("\n")
    chunks, current = [], ""
    for p in paras:
        if len(current) + len(p) > max_chars and current:
            chunks.append(current)
            current = p
        else:
            current += "\n" + p
    if current.strip():
        chunks.append(current)
    return chunks


# ---------------------------------------------------------------------------
# File readers
# ---------------------------------------------------------------------------

def read_docx(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Extraction per file
# ---------------------------------------------------------------------------

def extract_generic_from_text(text, filename, model):
    records = []
    for i, chunk in enumerate(chunk_text(text)):
        print(f"    chunk {i+1}")
        try:
            response = call_ollama(GENERIC_EXTRACTION_PROMPT.format(chunk=chunk), model)
        except requests.RequestException as e:
            print(f"    ! Ollama call failed: {e}", file=sys.stderr)
            continue
        for r in extract_json_blob(response, want_array=True):
            r["_file"] = filename
            records.append(r)
    return records


CHECKLIST_SUFFIXES = (".docx", ".doc")


def norm(s):
    if s is None:
        return None
    s = re.sub(r"\s+", " ", str(s)).strip().lower()
    return s or None


def drawing_from_record(rec: dict) -> dict:
    """sheetreader record -> the flat dict the rules below work on."""
    tb = rec["title_block"]
    return {
        "_file": rec["file"], "_page": rec["page"], "_ocr": rec["text_source"] == "ocr",
        "drawing_no": tb.get("drawing_number"), "revision": tb.get("revision"),
        "date": tb.get("date"), "title": tb.get("title"), "material": tb.get("material"),
        "weight": tb.get("weight"), "next_assy": tb.get("next_assy"), "sheet": tb.get("sheet"),
        "form_id": rec.get("form_id"), "general_tolerance": rec.get("general_tolerance") or {},
        "filename_drawing_no": rec.get("filename_drawing_number"),
        "bom": [{"item_no": r.get("item"), "description": r.get("description"), "qty": r.get("qty"),
                 "ref": r.get("ref") or r.get("part_number")} for r in rec.get("bom") or []],
    }


def canonicalize_ocr_ids(drawings):
    """On OCR'd sheets, replace drawing numbers that match a known number
    only up to look-alike characters ('AI-KD M12-E-02' -> 'AT-KD M12-E-02').
    Known numbers come from file names and text-layer sheets, which are exact."""
    known = {}
    for d in drawings:
        for dn in (d.get("filename_drawing_no"), None if d.get("_ocr") else d.get("drawing_no")):
            if dn:
                known.setdefault(ocr_key(dn), dn)
    for d in drawings:
        if not d.get("_ocr"):
            continue
        d["_ocr_corrections"] = []

        def fix(v):
            if not v or ocr_key(v) not in known or alnum(v) == alnum(known[ocr_key(v)]):
                return v
            d["_ocr_corrections"].append(f"{v} -> {known[ocr_key(v)]}")
            return known[ocr_key(v)]

        if d.get("drawing_no") and d.get("filename_drawing_no") and \
                alnum(d["drawing_no"]).replace("AND", "") != alnum(d["filename_drawing_no"]):
            d["drawing_no"] = fix(d["drawing_no"])
        d["next_assy"] = fix(d.get("next_assy"))
        for line in d.get("bom", []):
            line["ref"] = fix(line.get("ref"))


def _letter_prefix(s):
    m = re.match(r"[A-Z]{2,}", alnum(s))
    return m.group(0) if m else None


def mark_drawing_refs(drawings):
    """A BOM 'ref' is treated as a drawing reference when it starts like the
    drawing numbers in this set (e.g. 'AT-KD M12-E-04' in an AT-KD set), so
    hardware ('M2.5x4 CSK SCREWS'), materials ('ARALDITE 2011') and vendor
    part numbers ('P/N: PN1234') are not expected to have drawings."""
    prefixes = {_letter_prefix(d.get("drawing_no") or d.get("filename_drawing_no")) for d in drawings}
    prefixes.discard(None)
    for d in drawings:
        for line in d.get("bom", []):
            if "ref_part_no" in line:
                continue
            ref = line.get("ref")
            line["ref_part_no"] = ref if ref and _letter_prefix(ref) in prefixes else None


def id_keys(dn, ocr=False):
    """Keys a drawing number answers to: its letters+digits, plus each part
    of an 'A&B' / 'A/B' combined drawing ('AT-KD M12-E-03 A&B' answers to
    E-03A and E-03B). OCR'd numbers also answer to their look-alike key."""
    if not dn:
        return set()
    keys = {alnum(dn)}
    m = re.match(r"^(.*?\d)[\s_-]*([A-Z])\s*[&/,]\s*([A-Z])\s*$", dn.strip(), re.I)
    if m:
        base = alnum(m.group(1))
        keys |= {base + m.group(2).upper(), base + m.group(3).upper()}
    if ocr:
        keys |= {ocr_key(k) for k in list(keys)}
    return keys


class DrawingIndex:
    def __init__(self, drawings):
        self.sheets = drawings
        self.by_key = defaultdict(list)
        for d in drawings:
            dn = d.get("drawing_no") or d.get("filename_drawing_no")
            for k in id_keys(dn, d.get("_ocr")) | id_keys(d.get("filename_drawing_no")):
                self.by_key[k].append(d)

    def find(self, ref, ocr=False):
        found = []
        for k in id_keys(ref, True) if ocr else id_keys(ref):
            for d in self.by_key.get(k, []):
                if d not in found:
                    found.append(d)
        return found

    def find_any(self, ref):
        """Exact match first; failing that, allow OCR look-alike characters
        (only OCR'd sheets are indexed under look-alike keys)."""
        return self.find(ref) or self.find(ref, ocr=True)


def _issue(kind, severity, dwg, detail, sheets):
    return {"type": kind, "severity": severity, "drawing_no": dwg, "detail": detail, "sheets": sheets}


def _titles_agree(a, b):
    na, nb = alnum(a), alnum(b)
    return bool(na and nb) and (na == nb or na in nb or nb in na)


def check_drawings(drawings):
    """Cross-drawing rules. Severity: high = definite contradiction, medium =
    probable error, low = can't be checked from this folder."""
    canonicalize_ocr_ids(drawings)
    mark_drawing_refs(drawings)
    issues = []
    idx = DrawingIndex(drawings)

    by_no = defaultdict(list)
    for d in drawings:
        if d.get("drawing_no"):
            by_no[ocr_key(d["drawing_no"]) if d.get("_ocr") else alnum(d["drawing_no"])].append(d)
    # Same drawing on several sheets / files.
    for key, sheets in by_no.items():
        if len(sheets) < 2:
            continue
        for field, kind in (("revision", "Drawing revision mismatch"), ("date", "Drawing date mismatch"),
                            ("title", "Drawing title mismatch")):
            vals = {alnum(s.get(field)) for s in sheets if s.get(field)}
            if len(vals) > 1:
                issues.append(_issue(kind, "high", sheets[0]["drawing_no"],
                                     {s.get(field) for s in sheets if s.get(field)}, sheets))

    # BOM -> child drawings.
    for d in drawings:
        parent = d.get("drawing_no")
        for line in d.get("bom", []):
            ref = line.get("ref_part_no")
            if not ref:
                continue
            children = idx.find_any(ref)
            if not children:
                issues.append(_issue("BOM references missing drawing", "low", parent,
                                     f"BOM item {line.get('item_no') or '?'} '{line.get('description')}' "
                                     f"references {ref}, which is not in this folder - add it to check it.",
                                     [d]))
                continue
            for c in children:
                if line.get("description") and c.get("title") and not _titles_agree(line["description"], c["title"]):
                    issues.append(_issue("BOM description vs drawing title mismatch", "medium", parent,
                                         f"BOM calls {ref} '{line['description']}' but that drawing's title "
                                         f"is '{c['title']}'.", [d, c]))
                if c.get("next_assy") and parent and not (
                        id_keys(c["next_assy"], c.get("_ocr") or d.get("_ocr")) &
                        id_keys(parent, c.get("_ocr") or d.get("_ocr"))):
                    issues.append(_issue("Child's NEXT ASSY. is a different drawing", "medium", parent,
                                         f"{parent}'s BOM lists {c['drawing_no']}, but {c['drawing_no']}'s "
                                         f"NEXT ASSY. is {c['next_assy']}.", [d, c]))

    # NEXT ASSY. -> parent BOM.
    for d in drawings:
        this_no, nxt = d.get("drawing_no"), d.get("next_assy")
        if not (this_no and nxt):
            continue
        parents = idx.find_any(nxt)
        if not parents:
            issues.append(_issue("NEXT ASSY. points to unknown drawing", "low", this_no,
                                 f"NEXT ASSY. = {nxt}, which is not in this folder - add it to check the link.",
                                 [d]))
            continue
        mine = id_keys(this_no, True)
        listed = any(mine & id_keys(line.get("ref_part_no") or line.get("ref"), True)
                     for p in parents for line in p.get("bom", []))
        if not listed:
            issues.append(_issue("NEXT ASSY. not reciprocated in parent BOM", "high", this_no,
                                 f"{this_no} says NEXT ASSY. = {nxt}, but {nxt}'s BOM does not list {this_no}.",
                                 [d] + parents))

    # Same template form, different general tolerances.
    by_form = defaultdict(list)
    for d in drawings:
        if d.get("form_id") and d.get("general_tolerance"):
            by_form[re.sub(r"\s*REV.*$", "", d["form_id"])].append(d)
    for form, sheets in by_form.items():
        sigs = defaultdict(list)
        for s in sheets:
            gt = s["general_tolerance"]
            sig = tuple((k, gt.get(k)) for k in ("without_decimal", "one_decimal", "two_decimal") if gt.get(k))
            if sig:
                sigs[sig].append(s)
        if len(sigs) > 1:
            parts = [", ".join(f"{k.replace('_', ' ')} {v}" for k, v in sig) + " on " +
                     ", ".join(x["drawing_no"] or x["_file"] for x in ss) for sig, ss in sigs.items()]
            issues.append(_issue("Same template form, different general tolerances", "medium", form,
                                 f"Drawings on form {form} disagree: " + " | ".join(parts) +
                                 ". Confirm which is intended for each part.", sheets))
    return issues


# ---------------------------------------------------------------------------
# Dimension checklists
# ---------------------------------------------------------------------------

def read_checklist(path):
    """Header facts from a dimension checklist (material, weight), or None if
    the document isn't one (no zone/dimension table)."""
    from dimcheck.checklist_io import ensure_docx, load_checklist_rows
    try:
        document, table, rows = load_checklist_rows(str(path))
    except Exception:
        return None
    if not rows:
        return None
    text = "\n".join(c.text for r in table.rows for c in r.cells)
    mat = re.search(r"Material\s*:\s*(.+)", text, re.I)
    wt = re.search(r"Weight[^:\n]*:\s*(.+)", text, re.I)
    return {"_file": Path(path).name, "path": str(path),
            "material": mat.group(1).strip() if mat else None,
            "weight": wt.group(1).strip() if wt else None, "rows": len(rows)}


def _common_suffix(a, b):
    n = 0
    while n < min(len(a), len(b)) and a[-1 - n] == b[-1 - n]:
        n += 1
    return n


def match_checklist(chk, drawing_files):
    """The drawing a checklist belongs to: the longest shared file-name
    ending ('AT-F-DCL-PX_2B_40-03_REV._00' <-> 'AT-PX_2B_40-03_REV._00')."""
    stem = alnum(Path(chk["_file"]).stem)
    best, best_n = None, 0
    for f in drawing_files:
        n = _common_suffix(stem, alnum(Path(f).stem))
        if n > best_n:
            best, best_n = f, n
    return best if best_n >= 8 else None


def _weight_numbers(s):
    return [float(x) for x in re.findall(r"\d+(?:\.\d+)?", s or "")][:2]


def check_checklists(checklists, drawings, pdf_paths, run_dimcheck=True):
    issues = []
    for chk in checklists:
        pdf = match_checklist(chk, [p.name for p in pdf_paths])
        if not pdf:
            issues.append(_issue("Checklist without matching drawing", "low", chk["_file"],
                                 "No drawing in this folder matches this checklist's file name.", []))
            continue
        d = next((x for x in drawings if x["_file"] == pdf), None)
        if d is None:
            continue
        if chk.get("material") and d.get("material") and alnum(chk["material"]) != alnum(d["material"]):
            issues.append(_issue("Checklist material differs from drawing", "high", d["drawing_no"],
                                 f"Checklist: '{chk['material']}'; drawing: '{d['material']}'.", [d]))
        if chk.get("weight") and d.get("weight") and _weight_numbers(chk["weight"]) != _weight_numbers(d["weight"]):
            issues.append(_issue("Checklist weight differs from drawing", "high", d["drawing_no"],
                                 f"Checklist: '{chk['weight']}'; drawing: '{d['weight']}'.", [d]))
        if run_dimcheck:
            from dimcheck.pipeline import process
            path = next(p for p in pdf_paths if p.name == pdf)
            try:
                reports = process(str(path), chk["path"], None)
            except Exception as e:
                issues.append(_issue("Checklist could not be verified", "low", d["drawing_no"], str(e), [d]))
                continue
            chk["dimcheck"] = reports
            for r in reports:
                if r.verdict.startswith("MISMATCH"):
                    issues.append(_issue("Checklist tolerance differs from drawing", "high", d["drawing_no"],
                                         f"{chk['_file']}: zone {r.zone} dimension {r.dimension} - checklist "
                                         f"{r.existing_tolerance}, drawing {r.resolved_tolerance} ({r.source}).",
                                         [d]))
                elif r.verdict in ("not set", "kept") and r.source == "unresolved":
                    issues.append(_issue("Checklist row not verifiable", "low", d["drawing_no"],
                                         f"{chk['_file']}: zone {r.zone} dimension {r.dimension} - "
                                         + "; ".join(r.notes), [d]))
    return issues


# ---------------------------------------------------------------------------
# Generic (Word / Excel, model-extracted) cross-checks
# ---------------------------------------------------------------------------

def group_by_part_number(records):
    groups = defaultdict(list)
    for r in records:
        key = norm(r.get("part_number"))
        if key:
            groups[key].append(r)
    return groups


def find_generic_mismatches(groups):
    issues = []
    for pn, records in groups.items():
        names = {norm(r.get("part_name")) for r in records if r.get("part_name")}
        revs = {norm(r.get("revision")) for r in records if r.get("revision")}
        dates = {norm(r.get("date")) for r in records if r.get("date")}
        if len(names) > 1:
            issues.append({"type": "Part name mismatch", "part_number": pn, "detail": names, "records": records})
        if len(revs) > 1:
            issues.append({"type": "Revision mismatch", "part_number": pn, "detail": revs, "records": records})
        if len(dates) > 1:
            issues.append({"type": "Date mismatch", "part_number": pn, "detail": dates, "records": records})
    return issues


def fuzzy_cross_check(all_records, model):
    unmatched = [r for r in all_records if r.get("part_name") and not norm(r.get("part_number"))]
    suspects = []
    seen = set()
    for i, a in enumerate(unmatched):
        for b in unmatched[i + 1:]:
            if norm(a.get("part_name")) == norm(b.get("part_name")):
                continue
            if a["_file"] == b["_file"]:
                continue
            key = tuple(sorted([norm(a.get("part_name")), norm(b.get("part_name"))]))
            if key in seen:
                continue
            seen.add(key)
            prompt = FUZZY_MATCH_PROMPT.format(a=json.dumps(a), b=json.dumps(b))
            try:
                verdict = call_ollama(prompt, model).strip().upper()
            except requests.RequestException:
                continue
            if verdict.startswith("SAME"):
                suspects.append((a, b))
    return suspects


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

SEV_ORDER = {"high": 0, "medium": 1, "low": 2, "info": 3}


def write_report(drawing_issues, generic_issues, suspects, drawings, checklists, out_path, files_scanned):
    L = ["# Document Coherence Report", "", f"Files scanned: {len(files_scanned)}", ""]
    for f in files_scanned:
        L.append(f"- {f}")
    L += ["", "## Drawings read", "",
          "| File | Drawing No. | Rev | Date | Title | Next Assy. | BOM lines | Read via |",
          "|---|---|---|---|---|---|---|---|"]
    for d in drawings:
        L.append(f"| {d['_file']} | {d.get('drawing_no') or ''} | {d.get('revision') or ''} | {d.get('date') or ''} "
                 f"| {d.get('title') or ''} | {d.get('next_assy') or ''} | {len(d.get('bom', []))} "
                 f"| {'OCR' if d.get('_ocr') else 'text layer'} |")
    if checklists:
        L += ["", "## Checklists", ""]
        for c in checklists:
            reps = c.get("dimcheck") or []
            n_ok = sum(1 for r in reps if r.verdict == "match")
            L.append(f"- {c['_file']}: {c['rows']} dimension rows" +
                     (f", {n_ok} match the drawing" if reps else ""))
    L.append("")
    issues = sorted(drawing_issues, key=lambda i: SEV_ORDER.get(i.get("severity", "medium"), 9))
    if not issues and not generic_issues and not suspects:
        L.append("No inconsistencies found across files.")
    if issues:
        L += [f"## Cross-file issues: {len(issues)}", ""]
        for i in issues:
            L.append(f"- **[{i.get('severity', 'medium').upper()}] {i['type']}** - `{i['drawing_no']}`: {i['detail']}")
        L.append("")
    if generic_issues:
        L += [f"## Word/Excel mismatches: {len(generic_issues)}", ""]
        for issue in generic_issues:
            L.append(f"### {issue['type']} - part number `{issue['part_number']}`")
            L.append(f"Conflicting values: {', '.join(str(d) for d in issue['detail'])}")
            L.append("")
    if suspects:
        L += [f"## Possible same-part naming inconsistencies: {len(suspects)}", ""]
        for a, b in suspects:
            L.append(f"- `{a.get('part_name')}` ({a['_file']}) may be `{b.get('part_name')}` ({b['_file']}).")
    Path(out_path).write_text("\n".join(L), encoding="utf-8")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main(argv=None):
    from sheetreader.record import read_drawing

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder", help="Folder with drawing PDFs, checklists and other .docx/.xlsx files")
    ap.add_argument("--model", default=DEFAULT_MODEL)
    ap.add_argument("--out", default="coherence_report.md")
    ap.add_argument("--no-model", action="store_true",
                    help="Skip model extraction of generic Word/Excel files (drawings and checklists don't need it).")
    ap.add_argument("--skip-fuzzy", action="store_true")
    ap.add_argument("--no-dimcheck", action="store_true", help="Don't verify checklist tolerance rows.")
    args = ap.parse_args(argv)

    folder = Path(args.folder)
    files = sorted(p for p in folder.iterdir()
                   if p.suffix.lower() in (".docx", ".doc", ".xlsx", ".pdf") and not p.name.startswith("~$"))
    if not files:
        print("No .pdf/.docx/.doc/.xlsx files found.", file=sys.stderr)
        sys.exit(1)

    drawings, checklists, generic = [], [], []
    pdfs = [p for p in files if p.suffix.lower() == ".pdf"]
    for f in files:
        print(f"Reading {f.name} ...")
        if f.suffix.lower() == ".pdf":
            try:
                for rec in read_drawing(f):
                    drawings.append(drawing_from_record(rec.to_dict()))
            except Exception as e:
                print(f"  ! could not read {f.name}: {e}", file=sys.stderr)
            continue
        if f.suffix.lower() in CHECKLIST_SUFFIXES:
            chk = read_checklist(f)
            if chk:
                checklists.append(chk)
                continue
            if f.suffix.lower() == ".doc":
                continue
        if args.no_model:
            continue
        text = read_docx(f) if f.suffix.lower() == ".docx" else read_xlsx(f)
        if text.strip():
            generic.extend(extract_generic_from_text(text, f.name, args.model))

    drawing_issues = check_drawings(drawings)
    drawing_issues += check_checklists(checklists, drawings, pdfs, run_dimcheck=not args.no_dimcheck)
    generic_issues = find_generic_mismatches(group_by_part_number(generic))
    suspects = []
    if generic and not args.skip_fuzzy and not args.no_model:
        suspects = fuzzy_cross_check(generic, args.model)

    write_report(drawing_issues, generic_issues, suspects, drawings, checklists, args.out, [f.name for f in files])
    n = defaultdict(int)
    for i in drawing_issues:
        n[i.get("severity", "medium")] += 1
    print(f"\nReport written to {args.out}: {n['high']} high, {n['medium']} medium, {n['low']} low "
          f"cross-file issue(s); {len(generic_issues)} Word/Excel mismatch(es).")
    return drawing_issues


if __name__ == "__main__":
    main()
