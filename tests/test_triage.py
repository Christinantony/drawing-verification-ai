"""Triage: file inventory, content-based selection of what the supervisor
reviews, and warnings for everything that is skipped."""
import json
import shutil

import docx
import pytest

from supervisor import triage
from supervisor.__main__ import main
from supervisor.controller import Controller
from supervisor.store import Store
from supervisor.workspace import Workspace
from make_cad_like_drawing import make_cad_like_drawing

from test_supervisor import FakeModel


def _checklist(path, rows):
    d = docx.Document()
    t = d.add_table(rows=len(rows) + 1, cols=4)
    for c, h in zip(t.rows[0].cells, ("Sl No", "ZONE", "DIMENSIONS", "TOLERANCE")):
        c.text = h
    for i, (zone, dim, tol) in enumerate(rows, 1):
        cells = t.rows[i].cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = str(i), zone, dim, tol
    d.save(str(path))


def _docx(path, text):
    d = docx.Document()
    d.add_paragraph(text)
    d.save(str(path))


@pytest.fixture
def folder(tmp_path):
    f = tmp_path / "set"
    f.mkdir()
    return f


def _workspace(folder, tmp_path, review_all=False, clean=False):
    store = Store(tmp_path / "out" / "evidence.sqlite3")
    ws = Workspace(folder, store, tmp_path / "out")
    ws.prepare()
    if clean:
        # The synthetic sheet always carries "SILICON SEALANT"; drop that one
        # known-misspelling finding to model a sheet the rules pass.
        ws.baseline = [f for f in ws.baseline if "SILICON" not in f["claim"]]
    triage.plan(ws, review_all=review_all)
    return ws, store


GOOD_ROWS = [("C4", "50.0", "0/-0.2"), ("C3", "30.5", "±0.1"), ("B4", "40.0", "0/-0.1")]


def test_unsupported_files_are_listed_with_a_warning(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))
    (folder / "XY-100-01.dwg").write_bytes(b"AC1032")
    (folder / "notes.txt").write_text("hello")
    (folder / "Thumbs.db").write_bytes(b"x")
    ws, store = _workspace(folder, tmp_path)
    try:
        assert ws.plan["files"]["XY-100-01.dwg"] == "not supported"
        assert ws.plan["files"]["notes.txt"] == "not supported"
        assert "Thumbs.db" not in ws.plan["files"]
        msgs = [w["message"] for w in ws.warnings if w["code"] == "unsupported_file"]
        assert any("XY-100-01.dwg" in m and "export the sheet as PDF" in m for m in msgs)
    finally:
        store.close()


def test_folder_with_only_unsupported_files_says_what_it_found(folder, tmp_path):
    (folder / "part.step").write_bytes(b"ISO-10303")
    store = Store(tmp_path / "out" / "e.sqlite3")
    try:
        with pytest.raises(ValueError, match="part.step"):
            Workspace(folder, store, tmp_path / "out")
    finally:
        store.close()


def test_concerning_drawing_goes_to_the_supervisor_with_reasons(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))  # BRAKCET, SCALE 10:1, qty 0
    ws, store = _workspace(folder, tmp_path)
    try:
        drawings = [p for p in ws.plan["planned"] if p["kind"] == "drawing"]
        assert drawings
        reasons = " ".join(drawings[0]["reasons"])
        assert "rule finding" in reasons and "BRAKCET" in reasons
        assert {e for g in ws.groups for e in g["required_ids"]} >= set(drawings[0]["evidence_ids"])
    finally:
        store.close()


def test_clean_drawing_is_skipped_with_a_warning(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"), title="SUPPORT BRACKET", scale="1:1", qty3="1")
    ws, store = _workspace(folder, tmp_path, clean=True)
    try:
        assert not [p for p in ws.plan["planned"] if p["kind"] == "drawing"]
        assert any(s["kind"] == "drawing" for s in ws.plan["skipped"])
        assert any(w["code"] == "skipped_not_concerning" and "all rules passed" in w["message"] for w in ws.warnings)
        result = Controller(ws, FakeModel()).review()
        assert result["coverage"] == "complete"            # nothing planned was left undone
        assert result["skipped"] and result["warnings"]
    finally:
        store.close()


def test_all_flag_reviews_everything(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"), title="SUPPORT BRACKET", scale="1:1", qty3="1")
    ws, store = _workspace(folder, tmp_path, review_all=True, clean=True)
    try:
        assert not ws.plan["skipped"]
        assert any(p["kind"] == "drawing" for p in ws.plan["planned"])
    finally:
        store.close()


def test_checklist_without_its_drawing_is_skipped_with_a_warning(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))
    _checklist(folder / "XY-F-DCL-QQ-777-02_REV.03.docx", GOOD_ROWS)
    ws, store = _workspace(folder, tmp_path)
    try:
        assert ws.checklist_status["XY-F-DCL-QQ-777-02_REV.03.docx"] == "unmatched"
        assert any(w["code"] == "checklist_without_drawing" for w in ws.warnings)
        assert not any("Checklist" in u["question"] for u in ws.incomplete)
        assert any(s["reason"] == "no matching drawing in this set" for s in ws.plan["skipped"])
    finally:
        store.close()


def test_matching_checklist_is_skipped_mismatch_is_reviewed(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))
    _checklist(folder / "XY-F-DCL-XY-100-01_REV.01.docx", GOOD_ROWS + [("B3", "12", "")])
    ws, store = _workspace(folder, tmp_path)
    try:
        assert ws.checklist_status["XY-F-DCL-XY-100-01_REV.01.docx"] == "clean"
        assert not [p for p in ws.plan["planned"] if p["kind"] in ("checklist", "check_result")]
        # The empty cell is filled in the copy and reported, not an incomplete check.
        assert any(w["code"] == "tolerance_inserted" and "±1" in w["message"] for w in ws.warnings)
    finally:
        store.close()

    shutil.rmtree(tmp_path / "out")
    _checklist(folder / "XY-F-DCL-XY-100-01_REV.01.docx", [("C4", "50.0", "±0.5")] + GOOD_ROWS[1:])
    ws, store = _workspace(folder, tmp_path)
    try:
        assert ws.checklist_status["XY-F-DCL-XY-100-01_REV.01.docx"] == "concern"
        assert [p for p in ws.plan["planned"] if p["kind"] in ("checklist", "check_result")]
        assert any(f["severity"] == "high" and "C4" in f["claim"] for f in ws.baseline)
    finally:
        store.close()


def test_generic_document_reviewed_only_when_it_names_a_set_drawing(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))
    _docx(folder / "meeting-notes.docx", "Canteen menu for Friday: rice, dal, curd.")
    _docx(folder / "parts-list.docx", "XY-100-01 SUPPORT BRACKET Rev 01")
    ws, store = _workspace(folder, tmp_path)
    try:
        planned = {p["label"] for p in ws.plan["planned"] if p["kind"] == "document"}
        skipped = {s["label"] for s in ws.plan["skipped"] if s["kind"] == "document"}
        assert planned == {"parts-list.docx"} and skipped == {"meeting-notes.docx"}
    finally:
        store.close()


def test_plan_command_lists_files_and_uses_no_model(folder, tmp_path, capsys):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))
    (folder / "XY-100-01.dwg").write_bytes(b"AC1032")
    out = tmp_path / "out"
    assert main(["plan", str(folder), "--output", str(out), "--host", "http://127.0.0.1:9"]) == 0
    printed = capsys.readouterr().out
    assert "Supervisor (Gemma) review:" in printed and "No model was used" in printed
    md = next((out / "plans").glob("*.md")).read_text(encoding="utf-8")
    assert "XY-100-01.dwg" in md and "not supported" in md and "## Warnings" in md
    data = json.loads(next((out / "plans").glob("*.json")).read_text(encoding="utf-8"))
    assert data["plan"]["files"]["XY-100-01_REV.01.pdf"] == "read"


def _cells(path):
    t = docx.Document(str(path)).tables[0]
    return [[c.text for c in row.cells] for row in t.rows]


def _fill(path, row):
    from docx.oxml.ns import qn
    tc = docx.Document(str(path)).tables[0].rows[row].cells[3]._tc
    shd = tc.tcPr.find(qn("w:shd")) if tc.tcPr is not None else None
    return shd.get(qn("w:fill")) if shd is not None else None


def test_supervisor_fills_only_empty_tolerance_cells_in_a_copy(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))
    source = folder / "XY-F-DCL-XY-100-01_REV.01.docx"
    _checklist(source, [("C4", "50.0", ""),        # stacked limits on the drawing
                        ("C3", "30.5", "±0.5"),    # wrong value: flagged, never changed
                        ("B3", "12", ""),          # general table
                        ("A1", "999", "")])        # not on the drawing: stays empty
    before = _cells(source)
    ws, store = _workspace(folder, tmp_path)
    try:
        r = ws.checklist_results[source.name]
        copy = r["completed_copy"]
        assert copy and copy.endswith("XY-F-DCL-XY-100-01_REV.01_completed.docx")
        after = _cells(copy)
        assert _cells(source) == before                     # source untouched
        assert len(after) == len(before)                    # no rows added
        assert [row[3] for row in after[1:]] == ["0/-0.2", "±0.5", "±1", ""]
        assert [row[:3] for row in after] == [row[:3] for row in before]   # zones/dimensions unchanged
        assert _fill(copy, 1) == "FFF2CC" and _fill(copy, 3) == "FFF2CC"   # inserted, highlighted
        assert _fill(copy, 2) is None and _fill(copy, 4) is None          # not inserted, not highlighted
        assert r["inserted"] == 2 and r["mismatched"] == ["C3 30.5"] and r["still_empty"] == ["A1 999"]
        assert r["complete"] is False
        assert any("left empty" in u["question"] and "999" in u["question"] for u in ws.incomplete)
        assert any(f["severity"] == "high" and "C3" in f["claim"] for f in ws.baseline)
        result = Controller(ws, FakeModel()).review()
        assert result["checklists"][source.name]["inserted"] == 2
        from supervisor.report import write
        md = write(ws, result).read_text(encoding="utf-8")
        assert "## Checklists" in md and "A1 999" in md and "| NO |" in md
    finally:
        store.close()


def test_complete_checklist_needs_no_copy(folder, tmp_path):
    make_cad_like_drawing(str(folder / "XY-100-01_REV.01.pdf"))
    _checklist(folder / "XY-F-DCL-XY-100-01_REV.01.docx", GOOD_ROWS)
    ws, store = _workspace(folder, tmp_path)
    try:
        r = ws.checklist_results["XY-F-DCL-XY-100-01_REV.01.docx"]
        assert r["complete"] and r["inserted"] == 0 and r["completed_copy"] is None
        assert not list((tmp_path / "out").rglob("*_completed.docx"))
    finally:
        store.close()
