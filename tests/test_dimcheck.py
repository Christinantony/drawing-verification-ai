"""Tests for the dimcheck building blocks (zone grid, general tolerance
table, token extraction, dimension matching) and the full pipeline, run on
the synthetic drawing + a hand-built checklist .docx."""
import docx

from dimcheck.pdf_layout import load_page_layout, flat_text
from dimcheck.zoning import detect_grid
from dimcheck.general_tolerance import parse_general_tolerance_table, count_decimals
from dimcheck.dimension_matcher import find_dimension
from dimcheck.pipeline import resolve_tolerance, process
from make_synthetic_drawing import make_drawing


def _layout(tmp_path):
    pdf = tmp_path / "synthetic.pdf"
    make_drawing(str(pdf))
    return load_page_layout(str(pdf))


def test_zone_grid_detected_from_border_labels(tmp_path):
    layout = _layout(tmp_path)
    grid = detect_grid(layout.tokens, layout.page_width, layout.page_height)
    assert grid.col_labels == ["4", "3", "2", "1"]
    assert grid.row_labels == ["C", "B", "A"]
    assert grid.locate(300, 120) == "C3"
    assert grid.locate(540, 300) == "B2"
    assert not grid.warnings


def test_general_tolerance_table_parsed(tmp_path):
    general = parse_general_tolerance_table(flat_text(_layout(tmp_path)))
    assert general.without_decimal == "±1"
    assert general.one_decimal == "±0.1"
    assert general.two_decimal == "±0.05"
    assert general.angle == "±0.5°"
    assert general.lookup_by_decimals(count_decimals("25.0")) == "±0.1"


def test_count_decimals():
    assert count_decimals("12") == 0
    assert count_decimals("3.5") == 1
    assert count_decimals("0.25") == 2
    assert count_decimals("M3") is None


def test_find_dimension_locates_value_in_its_zone(tmp_path):
    layout = _layout(tmp_path)
    grid = detect_grid(layout.tokens, layout.page_width, layout.page_height)

    match = find_dimension(layout.tokens, grid, "C3", "25.0")
    assert match.matched_token is not None
    assert match.matched_token.text == "25.0"
    assert match.confidence == "high"
    # Nothing is printed next to this dimension on the synthetic drawing, so
    # neither an inline tolerance nor a fit code should be picked up.
    assert match.explicit_tolerance is None
    assert match.fit_code is None

    match = find_dimension(layout.tokens, grid, "B2", "12")
    assert match.matched_token is not None
    assert match.matched_token.text == "12"
    assert match.confidence == "high"


def test_find_dimension_not_found_is_low_confidence_with_a_note(tmp_path):
    layout = _layout(tmp_path)
    grid = detect_grid(layout.tokens, layout.page_width, layout.page_height)

    match = find_dimension(layout.tokens, grid, "A1", "999")
    assert match.matched_token is None
    assert match.confidence == "low"
    assert match.notes  # explains why, for the reviewer


def test_find_dimension_unknown_zone_is_reported_not_raised(tmp_path):
    layout = _layout(tmp_path)
    grid = detect_grid(layout.tokens, layout.page_width, layout.page_height)

    match = find_dimension(layout.tokens, grid, "Z9", "1.0")
    assert match.matched_token is None
    assert match.confidence == "low"
    assert "Z9" in match.notes[0]


def test_resolve_tolerance_falls_back_to_general_table_by_decimal_count(tmp_path):
    layout = _layout(tmp_path)
    grid = detect_grid(layout.tokens, layout.page_width, layout.page_height)
    general = parse_general_tolerance_table(flat_text(layout))

    value, source, _ = resolve_tolerance(find_dimension(layout.tokens, grid, "C3", "25.0"), general)
    assert value == "±0.1"
    assert source == "general_table"

    value, source, _ = resolve_tolerance(find_dimension(layout.tokens, grid, "B2", "12"), general)
    assert value == "±1"
    assert source == "general_table"


def _build_checklist(path, rows):
    """rows: list of (zone, dimension) tuples, left with an empty Tolerance cell."""
    document = docx.Document()
    table = document.add_table(rows=len(rows) + 1, cols=4)
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text, header[3].text = (
        "Sl No", "ZONE", "DIMENSIONS", "TOLERANCE",
    )
    for i, (zone, dimension) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = str(i), zone, dimension, ""
    document.save(str(path))


def test_pipeline_fills_checklist_tolerances_end_to_end(tmp_path):
    pdf_path = tmp_path / "synthetic.pdf"
    make_drawing(str(pdf_path))

    checklist_path = tmp_path / "checklist.docx"
    _build_checklist(checklist_path, [("C3", "25.0"), ("B2", "12")])

    output_path = tmp_path / "filled.docx"
    reports = process(str(pdf_path), str(checklist_path), str(output_path))

    by_zone = {r.zone: r for r in reports}
    assert by_zone["C3"].resolved_tolerance == "±0.1"
    assert by_zone["C3"].source == "general_table"
    assert by_zone["B2"].resolved_tolerance == "±1"
    assert by_zone["B2"].source == "general_table"

    filled_table = docx.Document(str(output_path)).tables[0]
    assert filled_table.rows[1].cells[3].text == "±0.1"
    assert filled_table.rows[2].cells[3].text == "±1"


def test_pipeline_leaves_unmatched_row_unresolved(tmp_path):
    pdf_path = tmp_path / "synthetic.pdf"
    make_drawing(str(pdf_path))

    checklist_path = tmp_path / "checklist.docx"
    _build_checklist(checklist_path, [("A1", "999")])  # not on the drawing

    output_path = tmp_path / "filled.docx"
    reports = process(str(pdf_path), str(checklist_path), str(output_path))

    assert reports[0].resolved_tolerance is None
    assert reports[0].source == "unresolved"
    assert reports[0].notes


# ---------------------------------------------------------------------------
# CAD-like drawing: stacked limits, "±" drawn as a graphic, vertical text,
# value matching, verify mode.
# ---------------------------------------------------------------------------

def _build_filled_checklist(path, rows):
    """rows: (zone, dimension, tolerance-or-'')"""
    document = docx.Document()
    table = document.add_table(rows=len(rows) + 1, cols=4)
    for c, h in zip(table.rows[0].cells, ("Sl No", "ZONE", "DIMENSIONS", "TOLERANCE")):
        c.text = h
    for i, (zone, dim, tol) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = str(i), zone, dim, tol
    document.save(str(path))


def test_stacked_symmetric_and_vertical_tolerances_are_read(cad_pdf):
    layout = load_page_layout(str(cad_pdf))
    grid = detect_grid(layout.tokens, layout.page_width, layout.page_height)
    stacked = find_dimension(layout.tokens, grid, "C4", "50.0")
    assert stacked.tolerance and stacked.tolerance.format() == "0/-0.2"
    sym = find_dimension(layout.tokens, grid, "C3", "30.5")
    assert sym.tolerance and sym.tolerance.format() == "±0.1"
    vertical = find_dimension(layout.tokens, grid, "B4", "40.0")
    assert vertical.tolerance and vertical.tolerance.format() == "0/-0.1"
    plain = find_dimension(layout.tokens, grid, "B3", "12")
    assert plain.tolerance is None and plain.matched_text == "12"


def test_checklist_values_match_drawing_by_value_not_text(cad_pdf):
    layout = load_page_layout(str(cad_pdf))
    grid = detect_grid(layout.tokens, layout.page_width, layout.page_height)
    m = find_dimension(layout.tokens, grid, "C4", "50")        # checklist "50", drawing "50.0"
    assert m.matched_text == "50.0"
    m = find_dimension(layout.tokens, grid, "C3", "Ø30.5")     # Ø is a graphic on the drawing
    assert m.matched_text == "30.5"


def test_verify_mode_reports_match_and_mismatch(cad_pdf, tmp_path):
    chk = tmp_path / "chk.docx"
    _build_filled_checklist(chk, [("C4", "50.0", "0/-0.2"), ("C3", "30.5", "±0.2"),
                                  ("B3", "12", ""), ("B4", "40.0", "0.0/-0.10")])
    out = tmp_path / "out.docx"
    reports = process(str(cad_pdf), str(chk), str(out))
    verdicts = {r.dimension: r.verdict for r in reports}
    assert verdicts == {"50.0": "match", "30.5": "MISMATCH", "12": "filled", "40.0": "match"}
    mism = next(r for r in reports if r.verdict == "MISMATCH")
    assert "Checklist says ±0.2, drawing says ±0.1" in mism.notes[0]
    table = docx.Document(str(out)).tables[0]
    assert table.rows[2].cells[3].text == "±0.2"   # existing value never overwritten by default
    assert table.rows[3].cells[3].text == "±1"     # empty cell filled from the general table


def test_tolerance_text_comparison():
    from dimcheck.pipeline import same_tolerance
    assert same_tolerance("0/-0.2", "+0.0/-0.20")
    assert same_tolerance("± 0.1", "±0.1")
    assert same_tolerance("-0.2/0", "0/-0.2")
    assert not same_tolerance("±0.1", "0/-0.1")
    assert same_tolerance("H7", "±0.1") is None
